"""Reading documents the company shared, and creating a case from them (no network, no LLM)."""
import io
import sys
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))

from navigator import api, cases, documents, llm  # noqa: E402

PITCH = "GlucoPatch develops a needle-free CGM patch for people with type 1 diabetes."
NEED = "We are looking for a clinical study partner in Switzerland."


def _docx() -> bytes:
    from docx import Document
    d = Document()
    d.add_heading("GlucoPatch", 0)
    d.add_paragraph(PITCH)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Stage", "Pre-clinical"
    b = io.BytesIO(); d.save(b); return b.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    p = Presentation()
    for title, body in [("GlucoPatch", PITCH), ("The ask", NEED)]:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text, s.placeholders[1].text = title, body
    b = io.BytesIO(); p.save(b); return b.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Company"
    ws.append(["Product", PITCH]); ws.append(["Need", NEED])
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def _pdf(text: str) -> bytes:
    """Smallest valid one-page PDF with a text line."""
    stream = f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offs.append(len(out)); out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % x for x in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


def _odt() -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", '<?xml version="1.0"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                   'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
                   f'<text:h>GlucoPatch</text:h><text:p>{PITCH}</text:p></office:text></office:body></office:document-content>')
    return b.getvalue()


@pytest.mark.parametrize("name,make,kind,pages", [
    ("pitch.docx", _docx, "Word", 0), ("deck.pptx", _pptx, "PowerPoint", 2), ("facts.xlsx", _xlsx, "Excel", 1),
    ("onepager.pdf", lambda: _pdf(PITCH), "PDF", 1), ("profile.odt", _odt, "OpenDocument text", 0),
    ("mail.txt", lambda: (PITCH + "\n" + NEED).encode(), "Text", 0),
])
def test_reads_common_formats(name, make, kind, pages):
    d = documents.read(name, make())
    assert d.kind == kind and d.pages == pages and "needle-free CGM" in d.text


def test_pptx_keeps_every_slide_and_tables_in_word():
    assert NEED in documents.read("deck.pptx", _pptx()).text
    assert "Stage | Pre-clinical" in documents.read("p.docx", _docx()).text


@pytest.mark.parametrize("name,data,msg", [
    ("old.doc", b"\xd0\xcf\x11\xe0" + b"x" * 100, "save it as PDF"),
    ("scan.pdf", _pdf(""), "scanned"),
    ("virus.exe", b"MZ" + b"x" * 100, "not supported"),
    ("fake.docx", b"not a zip at all, just text pretending" * 3, "damaged"),
    ("empty.pdf", b"", "empty"),
])
def test_clear_errors(name, data, msg):
    with pytest.raises(documents.DocumentError, match=msg):
        documents.read(name, data)


def test_zip_bomb_is_refused(monkeypatch):
    monkeypatch.setattr(documents, "MAX_UNZIPPED", 1000)
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("content.xml", "<a>" + "x" * 100_000 + "</a>")
    with pytest.raises(documents.DocumentError, match="too large"):
        documents.read("bomb.odt", b.getvalue())


def test_upload_then_case(tmp_path, monkeypatch):
    monkeypatch.setattr(cases, "DB", tmp_path / "cases.db")
    monkeypatch.setattr(llm, "available", lambda *a, **k: False)
    queries = []
    monkeypatch.setattr(cases, "find_matches", lambda a, k=12: (queries.append(a.search_query) or [], None))
    client = TestClient(api.app)

    doc = client.post("/documents?name=deck.pptx", content=_pptx()).json()
    assert doc["kind"] == "PowerPoint" and doc["pages"] == 2
    bad = client.post("/documents?name=old.ppt", content=b"\xd0\xcf\x11\xe0xxxx")
    assert bad.status_code == 422 and ".pptx" in bad.json()["detail"]
    monkeypatch.setattr(documents, "MAX_BYTES", 100)
    assert client.post("/documents?name=big.txt", content=b"x" * 500).status_code == 413
    monkeypatch.setattr(documents, "MAX_BYTES", 25_000_000)

    c = client.post("/cases", json={"documents": [doc], "company": "GlucoPatch", "actor": "Maria"}).json()
    assert c["documents"][0]["name"] == "deck.pptx" and "needle-free CGM" in c["analysis"]["summary"]
    assert "continuous glucose monitoring" in queries[0]
    assert any(e["kind"] == "document" for e in c["events"])


def test_archetypes_follow_the_inventory():
    from navigator.archetypes import classify
    assert classify("person", "Prof. Dr. A", "Group Head", "a@unibe.ch") == "expert"
    assert classify("person", "Dr. med. B", "Chief Physician", "") == "clinician"
    assert classify("person", "C", "Senior Researcher", "c@insel.ch") == "clinician"     # hospital email
    assert classify("person", "D", "MD-PhD Student", "d@insel.ch") == "workforce"        # learners stay workforce
    assert classify("person", "E", "Head of Regulatory and Quality Affairs") == "specialist"
    assert classify("person", "F", "Study Nurse") == "specialist"
    assert classify("person", "PD Dr. G", "Head Urogenital Engineering (UGE)") == "expert"
    assert classify("infrastructure", "ARTORG 3D-Printing Core Facility", text="rapid prototyping") == "facility"
    assert classify("infrastructure", "Clinical Research Database") == "digital"
    assert classify("infrastructure", "Biomechanical Testing") == "service"
    assert classify("infrastructure", "Unitectra") == "technology"
    assert classify("project", "SPN-005 – Pilot study on painless glucose monitoring") == "study"
    assert classify("project", "Master's Thesis") == "project"
