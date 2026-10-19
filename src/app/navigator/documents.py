"""Read a document the company shared (pitch deck, brochure, email, spreadsheet) into plain text.

Only the extracted text is kept, with the case; the uploaded file itself is never stored.
Uploads are untrusted: size-capped, zip archives (docx/pptx/xlsx/odt) checked against
decompression bombs before parsing, and page/slide/row counts bounded.
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import PurePath
from xml.etree import ElementTree

from pydantic import BaseModel, Field

from .crawl import page_text

MAX_BYTES = 25_000_000          # per upload
MAX_UNZIPPED = 250_000_000      # docx/pptx/xlsx/odf are zip archives
MAX_PAGES = 300
MAX_TEXT = 40_000               # kept per document

KINDS = {
    ".pdf": "PDF", ".docx": "Word", ".docm": "Word", ".pptx": "PowerPoint", ".pptm": "PowerPoint",
    ".xlsx": "Excel", ".xlsm": "Excel", ".odt": "OpenDocument text", ".odp": "OpenDocument slides",
    ".ods": "OpenDocument spreadsheet", ".rtf": "RTF", ".txt": "Text", ".md": "Text", ".csv": "CSV",
    ".html": "Web page", ".htm": "Web page", ".eml": "Email",
}
_LEGACY = {".doc": "Word", ".ppt": "PowerPoint", ".xls": "Excel", ".key": "Keynote", ".pages": "Pages"}
ACCEPT = ",".join(KINDS)


class DocumentError(ValueError):
    """Shown to the officer as is."""


class Document(BaseModel):
    original_id: str | None = Field(default=None, pattern=r'^[0-9a-f]{32}$')
    original_size: int = Field(default=0, ge=0, le=25_000_000)
    original_sha256: str | None = Field(default=None, pattern=r'^[0-9a-f]{64}$')
    name: str = Field(..., max_length=200)
    kind: str = Field(..., max_length=40)
    pages: int = 0               # pages / slides / sheets, when the format has them
    text: str = Field(..., max_length=MAX_TEXT)


def _clean(text: str) -> str:
    lines = [re.sub(r"[ \t ]+", " ", x).strip() for x in text.replace("\r", "\n").splitlines()]
    out: list[str] = []
    for x in lines:
        if x or (out and out[-1]):
            out.append(x)
    return "\n".join(out).strip()


def _zip(data: bytes) -> zipfile.ZipFile:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise DocumentError("The file is damaged or not really in the format its name suggests.") from None
    if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
        raise DocumentError("The file is too large once unpacked.")
    return z


def _pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        r = PdfReader(io.BytesIO(data))
        if r.is_encrypted and not r.decrypt(""):
            raise DocumentError("The PDF is password-protected. Please ask the company for an unprotected copy.")
        pages = r.pages[:MAX_PAGES]
        return "\n\n".join((p.extract_text() or "") for p in pages), len(r.pages)
    except PdfReadError:
        raise DocumentError("The PDF could not be read (it may be damaged).") from None


def _docx(data: bytes) -> tuple[str, int]:
    from docx import Document as Docx
    _zip(data)                       # size check before parsing
    d = Docx(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip())))
    return "\n".join(parts), 0


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation
    _zip(data)
    deck = Presentation(io.BytesIO(data))
    slides = []
    for i, s in enumerate(deck.slides):
        if i >= MAX_PAGES:
            break
        parts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            parts.append(s.notes_slide.notes_text_frame.text)
        slides.append("\n".join(p for p in parts if p.strip()))
    return "\n\n".join(slides), len(deck.slides)


def _xlsx(data: bytes) -> tuple[str, int]:
    from openpyxl import load_workbook
    _zip(data)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[{ws.title}]")
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= 2000:
                break
            cells = [str(v) for v in row if v not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out), len(wb.worksheets)


def _odf(data: bytes) -> tuple[str, int]:
    z = _zip(data)
    try:
        root = ElementTree.fromstring(z.read("content.xml"))
    except (KeyError, ElementTree.ParseError):
        raise DocumentError("The OpenDocument file could not be read.") from None
    blocks = []
    for el in root.iter():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag in ("p", "h"):
            blocks.append("".join(el.itertext()))
    pages = sum(1 for el in root.iter() if el.tag.endswith("}page") or el.tag.endswith("}table"))
    return "\n".join(blocks), pages


def _text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _eml(data: bytes) -> str:
    from email import policy
    from email.parser import BytesParser
    msg = BytesParser(policy=policy.default).parsebytes(data)
    head = "\n".join(f"{k}: {msg[k]}" for k in ("From", "Subject", "Date") if msg[k])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body and body.get_content_type() == "text/html":
        text = page_text(text)
    return f"{head}\n\n{text}"


def read(name: str, data: bytes) -> Document:
    name = PurePath(name or "document").name[:200]
    ext = PurePath(name.lower()).suffix
    if not data:
        raise DocumentError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError(f"The file is larger than {MAX_BYTES // 1_000_000} MB.")
    if ext in _LEGACY:
        raise DocumentError(f"Old {_LEGACY[ext]} files ({ext}) cannot be read. Please save it as PDF "
                            f"(or {'.docx' if ext == '.doc' else '.pptx' if ext == '.ppt' else '.xlsx' if ext == '.xls' else 'PDF'}) "
                            "and upload that.")
    if ext not in KINDS:
        raise DocumentError(f"“{ext or name}” files are not supported. Supported: PDF, Word, PowerPoint, Excel, "
                            "OpenDocument, RTF, text, HTML and email (.eml).")
    pages = 0
    try:
        if ext == ".pdf":
            text, pages = _pdf(data)
        elif ext in (".docx", ".docm"):
            text, pages = _docx(data)
        elif ext in (".pptx", ".pptm"):
            text, pages = _pptx(data)
        elif ext in (".xlsx", ".xlsm"):
            text, pages = _xlsx(data)
        elif ext in (".odt", ".odp", ".ods"):
            text, pages = _odf(data)
        elif ext == ".rtf":
            from striprtf.striprtf import rtf_to_text
            text = rtf_to_text(_text(data), errors="ignore")
        elif ext in (".html", ".htm"):
            text = page_text(_text(data))
        elif ext == ".eml":
            text = _eml(data)
        else:
            text = _text(data)
    except DocumentError:
        raise
    except Exception as e:  # parsers raise many different errors on damaged files
        raise DocumentError(f"The file could not be read ({type(e).__name__}). It may be damaged; "
                            "try saving it again as PDF.") from None
    text = _clean(text)
    if len(text) < 40:
        hint = (" It looks like a scanned document or slides made of images, which contain no text to read."
                if ext in (".pdf", ".pptx", ".pptm", ".odp") else "")
        raise DocumentError("No readable text was found in the file." + hint +
                            " Please paste the key information into the text box instead.")
    return Document(name=name, kind=KINDS[ext], pages=pages, text=text[:MAX_TEXT])
